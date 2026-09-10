"""Builds docs/readme.docx, the readme that is actually submitted, from README.md.

Run it after changing README.md, then run package.bat:

    python docs/make-readme-docx.py <identity-number>

The Word file is the one place the identity number appears. README.md lives in a public
repository and must never carry it, which is why the number is passed in as an argument
rather than stored anywhere. docs/readme.docx is gitignored for the same reason.

Only the small slice of Markdown the readme actually uses is handled: headings, tables,
fenced code, bullet and numbered lists, block quotes, horizontal rules and bold spans.
"""

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

PROJECT = Path(__file__).resolve().parents[1]
SOURCE = PROJECT / "README.md"
TARGET = PROJECT / "docs" / "readme.docx"

# A note addressed to whoever reads the public repository. In the Word file, which is the one
# place the number actually appears, it would only read as a contradiction.
NOTE_FOR_THE_REPOSITORY_ONLY = "תעודת הזהות מופיעה"


def rightToLeft(paragraph):
    """Marks a paragraph as right to left, which Word needs stated explicitly for Hebrew."""
    properties = paragraph._p.get_or_add_pPr()
    properties.append(properties.makeelement(qn("w:bidi"), {}))
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    return paragraph


def withBoldSpans(paragraph, text):
    """Writes text into a paragraph, turning **these** into bold runs."""
    for index, piece in enumerate(re.split(r"\*\*(.+?)\*\*", text)):
        if piece:
            run = paragraph.add_run(piece.replace("`", ""))
            run.bold = index % 2 == 1
    return paragraph


def addTable(document, rows):
    header, *body = rows
    table = document.add_table(rows=1, cols=len(header))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.RIGHT
    for cell, caption in zip(table.rows[0].cells, header):
        withBoldSpans(rightToLeft(cell.paragraphs[0]), caption)
    for line in body:
        for cell, value in zip(table.add_row().cells, line):
            withBoldSpans(rightToLeft(cell.paragraphs[0]), value)
    document.add_paragraph()


def splitRow(line):
    return [part.strip() for part in line.strip().strip("|").split("|")]


def startsSomethingNew(stripped):
    """Whether a line begins a new block rather than continuing the one before it."""
    return (stripped.startswith(("#", "|", "```", "---", "> ", "- "))
            or re.match(r"^\d+\.\s", stripped) is not None)


def gatherWrappedLines(source, index):
    """README.md is hard wrapped, and a wrapped line is not a new paragraph."""
    block = []
    while index < len(source) and source[index].strip() and not startsSomethingNew(source[index].strip()):
        block.append(source[index].strip())
        index += 1
    return block, index


def looksLikeATable(source, index):
    if not source[index].strip().startswith("|") or index + 1 >= len(source):
        return False
    divider = source[index + 1].strip()
    return divider.startswith("|") and set(divider.replace("|", "").replace(":", "").strip()) <= {"-", " "}


def build(identityNumber):
    source = SOURCE.read_text(encoding="utf-8").splitlines()
    document = Document()
    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10.5)

    index = 0
    while index < len(source):
        stripped = source[index].strip()

        if stripped.startswith("```"):
            block = []
            index += 1
            while index < len(source) and not source[index].strip().startswith("```"):
                block.append(source[index])
                index += 1
            run = document.add_paragraph().add_run("\n".join(block))
            run.font.name = "Consolas"
            run.font.size = Pt(9.5)
            run.font.color.rgb = RGBColor(0x1F, 0x38, 0x64)
            index += 1
            continue

        if looksLikeATable(source, index):
            rows = [splitRow(stripped)]
            index += 2
            while index < len(source) and source[index].strip().startswith("|"):
                rows.append(splitRow(source[index]))
                index += 1
            addTable(document, rows)
            continue

        if NOTE_FOR_THE_REPOSITORY_ONLY in stripped:
            index += 1
            continue

        if stripped.startswith("---"):
            rightToLeft(document.add_paragraph()).add_run("─" * 40)
        elif stripped.startswith("### "):
            rightToLeft(document.add_heading(stripped[4:], level=3))
        elif stripped.startswith("## "):
            rightToLeft(document.add_heading(stripped[3:], level=2))
        elif stripped.startswith("# "):
            rightToLeft(document.add_heading(stripped[2:], level=1))
        elif stripped.startswith("> "):
            withBoldSpans(rightToLeft(document.add_paragraph()), stripped[2:]).runs[0].italic = True
        elif re.match(r"^\d+\.\s", stripped) or stripped.startswith("- "):
            numbered = not stripped.startswith("- ")
            first = re.sub(r"^\d+\.\s", "", stripped) if numbered else stripped[2:]
            rest, index = gatherWrappedLines(source, index + 1)
            style = "List Number" if numbered else "List Bullet"
            withBoldSpans(rightToLeft(document.add_paragraph(style=style)), " ".join([first] + rest))
            continue
        elif stripped:
            rest, index = gatherWrappedLines(source, index + 1)
            withBoldSpans(rightToLeft(document.add_paragraph()), " ".join([stripped] + rest))
            continue
        index += 1

    stampIdentityNumber(document, identityNumber)
    document.save(TARGET)
    print("wrote", TARGET)


def stampIdentityNumber(document, identityNumber):
    """Adds the identity number to the submitter table, which only the Word file carries."""
    for table in document.tables:
        for row in table.rows:
            if row.cells[0].text.strip() == "שם מלא":
                added = table.add_row().cells
                withBoldSpans(rightToLeft(added[0].paragraphs[0]), "ת.ז.")
                withBoldSpans(rightToLeft(added[1].paragraphs[0]), identityNumber)
                return
    raise SystemExit("Could not find the submitter table to add the identity number to.")


if __name__ == "__main__":
    if len(sys.argv) != 2 or not sys.argv[1].isdigit():
        raise SystemExit("usage: python docs/make-readme-docx.py <identity-number>")
    build(sys.argv[1])
